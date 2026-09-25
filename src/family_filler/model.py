import openpyxl

from ..core.logs import logger
from ..core.selectors import (
    HEADER_FILL_COLOR,
    HEADER_FONT_COLOR,
    REPORT_SHEETS,
    SHEET_NOT_ENTERED,
)
from ..core.settings import (
    EXCEL_DATE_MAX,
    EXCEL_DATE_MIN,
    NOTICE_TEXT_LIMIT,
    PINFL_LENGTH,
)
from ..core.utils import (
    birth_from_pinfl,
    digits_only,
    generate_uzbek_phone,
    norm_text,
    normalize_date,
)
from .config import (
    DOC_TYPE_DEFAULT,
    DOC_TYPE_MAP,
    HEADER_ROW,
    MAX_ROWS,
    NOT_ENTERED_HEADERS,
    NOT_ENTERED_WIDTHS,
    REPORT_HEADERS,
    REPORT_WIDTHS,
    START_ROW,
)
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from openpyxl.styles import Alignment, Font, PatternFill
from pathlib import Path
from typing import Dict, List, Optional, Tuple


@dataclass
class ExcelRow:
    """Строка family_members.xlsx: A № | B ID | C Ф.И.Ш. | D ДР | E Пол | F Национальность |
    G Гражданство | H ПИНФЛ | I Тип документа | J Серия | K Номер | L Организация |
    M ИНН | N Регион | O Класс."""
    row_index: int
    number: str = ""
    pid: str = ""
    full_name: str = ""
    birth: str = ""
    gender: str = ""
    nationality: str = ""
    citizenship: str = ""
    pinfl: str = ""
    doc_type: str = ""
    doc_series: str = ""
    doc_number: str = ""
    organization: str = ""
    org_inn: str = ""
    org_region: str = ""
    grade: str = ""
    phone: str = ""

    @property
    def code(self) -> str:
        """Ключ прогресса/отчёта — ID из таблицы (или номер строки)."""
        return self.pid or f"row{self.row_index}"

    def pinfl_digits(self) -> str:
        digits = digits_only(self.pinfl)
        return digits[:PINFL_LENGTH] if len(digits) >= PINFL_LENGTH else digits

    def birth_date(self) -> str:
        """Дата рождения: из таблицы, иначе вычисляется из ПИНФЛ."""
        return normalize_date(self.birth) or birth_from_pinfl(self.pinfl_digits()) or ""

    def doc_type_needles(self) -> Tuple[str, ...]:
        """Варианты названия типа документа для справочника сайта."""
        low = norm_text(self.doc_type)
        for key, needles in DOC_TYPE_MAP.items():
            if key in low:
                return needles
        return DOC_TYPE_DEFAULT

    def has_documents(self) -> bool:
        """Есть чем искать человека в модалке (серия/номер и дата рождения)."""
        return bool(self.doc_number and self.birth_date())

    def label(self) -> str:
        name = self.full_name or "—"
        doc = " ".join(x for x in (self.doc_series, self.doc_number) if x) or "—"
        return f"{name} | {doc} | {self.birth_date() or '—'}"

    def get_formatted_phone(self) -> str:
        return (self.phone or "").strip()


class ExcelReader:
    """Читает family_members.xlsx: заголовок в строке 6, данные с 7-й."""

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Файл не найден: {self.file_path}")

    def read_rows(self, start_row: int = START_ROW) -> List[ExcelRow]:
        workbook = openpyxl.load_workbook(self.file_path, data_only=True, read_only=True)
        worksheet = workbook.active
        rows: List[ExcelRow] = []
        seen = set()
        duplicates = 0
        empty_streak = 0
        first = max(HEADER_ROW + 1, start_row)
        skipped = first - (HEADER_ROW + 1)
        for idx, values in enumerate(
                worksheet.iter_rows(min_row=first, values_only=True), first):
            values = list(values) + [None] * (15 - len(values))
            row = ExcelRow(
                row_index=idx,
                number=self._value(values[0]),
                pid=self._value(values[1]),
                full_name=self._value(values[2]),
                birth=self._value(values[3]),
                gender=self._value(values[4]),
                nationality=self._value(values[5]),
                citizenship=self._value(values[6]),
                pinfl=self._value(values[7]),
                doc_type=self._value(values[8]),
                doc_series=self._value(values[9]),
                doc_number=self._value(values[10]),
                organization=self._value(values[11]),
                org_inn=self._value(values[12]),
                org_region=self._value(values[13]),
                grade=self._value(values[14]),
            )
            if not (row.full_name or row.pid or row.doc_number):
                empty_streak += 1
                if empty_streak > 500:
                    logger.info(f"Excel: пусто с строки {idx - 499} — чтение остановлено")
                    break
                continue
            empty_streak = 0
            if not row.has_documents():
                continue
            if row.code in seen:
                duplicates += 1
                continue
            seen.add(row.code)
            row.phone = generate_uzbek_phone()
            rows.append(row)
            if len(rows) >= MAX_ROWS:
                logger.warning(f"Excel: достигнут предел {MAX_ROWS} строк — дальше не читаем")
                break
        try:
            workbook.close()
        except Exception:
            pass
        logger.info(f"Excel: строк {len(rows)}, старт со строки листа {first}"
                    + (f", пропущено строк {skipped}" if skipped else "")
                    + (f", дубликатов пропущено {duplicates}" if duplicates else ""))
        return rows

    @staticmethod
    def _value(value) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%d.%m.%Y")
        if isinstance(value, str):
            text = value.strip()
            return "" if text.startswith("=") else text
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, int) and EXCEL_DATE_MIN < value < EXCEL_DATE_MAX:
            return (datetime(1899, 12, 30) + timedelta(days=value)).strftime("%d.%m.%Y")
        return str(value).strip()


class ReportWriter:
    def __init__(self, path: str):
        self.path = Path(path)
        if self.path.exists():
            try:
                self.wb = openpyxl.load_workbook(self.path)
            except Exception as exc:
                logger.warning(f"Отчёт не открылся ({exc}), создан новый")
                self.wb = self._new_workbook()
        else:
            self.wb = self._new_workbook()

        for name in REPORT_SHEETS:
            if name not in self.wb.sheetnames:
                self._create_sheet(name)
        self.save()

    @staticmethod
    def _new_workbook():
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        return wb

    def _create_sheet(self, name: str) -> None:
        ws = self.wb.create_sheet(name)
        headers = NOT_ENTERED_HEADERS if name == SHEET_NOT_ENTERED else REPORT_HEADERS
        widths = NOT_ENTERED_WIDTHS if name == SHEET_NOT_ENTERED else REPORT_WIDTHS
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for idx, width in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = width
        ws.freeze_panes = "A2"

    def add(self, sheet: str, street: str, code: str, full_name: str, series: str,
            number: str, birth: str, status: str, error_type: str = "",
            notice: str = "", url: str = "") -> None:
        ws = self.wb[sheet]
        ws.append([ws.max_row, datetime.now().strftime("%d.%m.%Y %H:%M:%S"),
                   street, code, full_name, series, number, birth, status,
                   error_type, (notice or "")[:NOTICE_TEXT_LIMIT], url])
        ws.cell(ws.max_row, 11).alignment = Alignment(wrap_text=True, vertical="top")
        self.save()

    def add_not_entered(self, street: str, code: str, full_name: str, series: str,
                        number: str, birth: str, error: str) -> None:
        ws = self.wb[SHEET_NOT_ENTERED]
        ws.append([ws.max_row, datetime.now().strftime("%d.%m.%Y %H:%M:%S"),
                   street, code, full_name, series, number, birth,
                   (error or "")[:NOTICE_TEXT_LIMIT]])
        ws.cell(ws.max_row, 9).alignment = Alignment(wrap_text=True, vertical="top")
        self.save()

    def save(self) -> None:
        try:
            self.wb.save(self.path)
        except PermissionError:
            alt = self.path.with_name(
                f"{self.path.stem}_{datetime.now().strftime('%H%M%S')}{self.path.suffix}")
            try:
                self.wb.save(alt)
                logger.warning(f"Отчёт занят, сохранён как {alt}")
                self.path = alt
            except Exception as exc:
                logger.error(f"Отчёт не сохранён: {exc}")
        except Exception as exc:
            logger.error(f"Отчёт не сохранён: {exc}")
